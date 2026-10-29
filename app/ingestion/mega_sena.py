from __future__ import annotations

import hashlib
import re
import unicodedata
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Iterable, Iterator, Mapping

from app.domain import (
    ComponentDefinition,
    ComponentResult,
    DrawCandidate,
    DrawStatus,
    ExternalReference,
    Lottery,
    LotteryRuleVersion,
    LotteryStatus,
    MonetaryAmount,
    OrderSemantics,
    PrizeTierResult,
    RawProvenance,
    ResultSeries,
    ResultSeriesDefinition,
    SeriesRole,
    ValidationIssue,
    ValueStatus,
    ValueType,
    validate_draw,
)


SOURCE_KEY = "caixa_mega_sena_xlsx"
RULE_VERSION = "baseline_1996"


MEGA_SENA = Lottery(
    slug="mega_sena",
    display_name="Mega-Sena",
    jurisdiction_code="BR",
    default_timezone="America/Sao_Paulo",
    default_currency="BRL",
    status=LotteryStatus.ACTIVE,
    rules=(
        LotteryRuleVersion(
            version_label=RULE_VERSION,
            valid_from=date(1996, 3, 11),
            source_reference="SPEC-0001: limites observados no arquivo histórico",
            series=(
                ResultSeriesDefinition(
                    key="primary",
                    label="Resultado principal",
                    role=SeriesRole.PRIMARY,
                    components=(
                        ComponentDefinition(
                            key="main",
                            label="Números",
                            value_type=ValueType.INTEGER,
                            min_count=6,
                            max_count=6,
                            min_value=1,
                            max_value=60,
                            unique_values=True,
                            order_semantics=OrderSemantics.SORTED,
                        ),
                    ),
                ),
            ),
        ),
    ),
)


@dataclass(frozen=True, slots=True)
class NormalizationResult:
    candidate: DrawCandidate | None
    issues: tuple[ValidationIssue, ...]
    canonical_key: str | None

    @property
    def is_valid(self) -> bool:
        return self.candidate is not None and not self.issues


def file_sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def iter_xlsx_rows(path: Path, *, limit: int | None = None) -> Iterator[dict[str, object]]:
    """Lê a primeira planilha em modo somente leitura e devolve linhas brutas."""
    if limit is not None and limit < 0:
        raise ValueError("limit não pode ser negativo")

    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("A leitura XLSX requer a dependência openpyxl") from exc

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        worksheet = workbook.worksheets[0]
        # A exportação atual declara incorretamente dimensão A1:T1, embora possua
        # milhares de linhas. No modo read-only, o openpyxl confia nesse metadado.
        worksheet.reset_dimensions()
        rows = worksheet.iter_rows(values_only=True)
        headers = next(rows, None)
        if not headers:
            raise ValueError("planilha não possui cabeçalho")
        header_names = tuple(str(value).strip() if value is not None else "" for value in headers)
        if "concurso" not in {_header_key(name) for name in header_names}:
            raise ValueError("cabeçalho Concurso não encontrado")

        produced = 0
        for values in rows:
            if limit is not None and produced >= limit:
                break
            if not any(value is not None for value in values):
                continue
            produced += 1
            yield dict(zip(header_names, values))
    finally:
        workbook.close()


def normalize_row(row: Mapping[str, object], provenance: RawProvenance) -> NormalizationResult:
    indexed = {_header_key(str(key)): value for key, value in row.items()}
    issues: list[ValidationIssue] = []

    draw_number = _parse_integer(
        _lookup(indexed, "concurso"), "draw.official_draw_number", issues
    )
    draw_date = _parse_date(
        _lookup(indexed, "datadosorteio", "datasorteio"), "draw.draw_date", issues
    )

    balls: list[int | str] = []
    for position in range(1, 7):
        path = f"draw.series.primary.components.main[{position - 1}]"
        raw_value = _lookup(indexed, f"bola{position}")
        value = _parse_integer(raw_value, path, issues)
        if value is not None:
            balls.append(value)

    if draw_number is None or draw_date is None:
        return NormalizationResult(None, tuple(issues), None)

    monetary_amounts = tuple(
        _parse_money(_lookup(indexed, *aliases), kind, f"draw.money.{kind}", issues)
        for kind, aliases in (
            ("rollover_amount", ("acumulado6acertos",)),
            ("sales_total", ("arrecadacaototal", "arrecadaototal")),
            ("next_estimated_jackpot", ("estimativapremio", "estimativaprmio")),
            (
                "special_draw_rollover",
                ("acumuladosorteioespecialmegadavirada",),
            ),
        )
    )

    prize_tiers = tuple(
        _parse_prize_tier(indexed, tier_key, hits, issues)
        for tier_key, hits in (("sena", 6), ("quina", 5), ("quadra", 4))
    )

    metadata = []
    city = _lookup(indexed, "cidadeuf")
    observation = _lookup(indexed, "observacao", "observao")
    if not _is_missing(city):
        metadata.append(("winner_locations_source_text", str(city).strip()))
    if not _is_missing(observation):
        metadata.append(("observation_source_text", str(observation).strip()))

    candidate = DrawCandidate(
        lottery_slug=MEGA_SENA.slug,
        draw_date=draw_date,
        rule_version_label=RULE_VERSION,
        official_draw_number=str(draw_number),
        status=DrawStatus.CANDIDATE,
        series=(
            ResultSeries(
                key="primary",
                components=(ComponentResult("main", tuple(balls)),),
            ),
        ),
        external_references=(
            ExternalReference(SOURCE_KEY, "official_draw_number", str(draw_number)),
        ),
        monetary_amounts=monetary_amounts,
        prize_tiers=prize_tiers,
        provenance=provenance,
        metadata=tuple(metadata),
    )
    issues.extend(validate_draw(candidate, MEGA_SENA))
    return NormalizationResult(
        candidate=candidate,
        issues=tuple(issues),
        canonical_key=f"mega_sena:official_draw_number:{draw_number}",
    )


def normalize_rows(
    rows: Iterable[Mapping[str, object]], provenance: RawProvenance
) -> tuple[NormalizationResult, ...]:
    return tuple(normalize_row(row, provenance) for row in rows)


def _header_key(value: str) -> str:
    decomposed = unicodedata.normalize("NFKD", value)
    ascii_text = "".join(character for character in decomposed if character.isascii())
    return re.sub(r"[^a-z0-9]", "", ascii_text.lower())


def _lookup(indexed: Mapping[str, object], *aliases: str) -> object | None:
    for alias in aliases:
        if alias in indexed:
            return indexed[alias]
    return None


def _is_missing(value: object) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def _parse_integer(
    value: object, path: str, issues: list[ValidationIssue]
) -> int | None:
    if _is_missing(value):
        issues.append(ValidationIssue("source.missing_value", path, "valor ausente"))
        return None
    if isinstance(value, bool):
        issues.append(
            ValidationIssue("source.invalid_integer", path, "valor não é inteiro", value, "integer")
        )
        return None
    if isinstance(value, int):
        return value
    if isinstance(value, float) and value.is_integer():
        return int(value)
    if isinstance(value, str) and re.fullmatch(r"[+-]?\d+", value.strip()):
        return int(value.strip())
    issues.append(
        ValidationIssue("source.invalid_integer", path, "valor não é inteiro", value, "integer")
    )
    return None


def _parse_date(
    value: object, path: str, issues: list[ValidationIssue]
) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str):
        try:
            return datetime.strptime(value.strip(), "%d/%m/%Y").date()
        except ValueError:
            pass
    issues.append(
        ValidationIssue("source.invalid_date", path, "data ausente ou inválida", value, "DD/MM/YYYY")
    )
    return None


def _parse_money(
    value: object,
    kind: str,
    path: str,
    issues: list[ValidationIssue],
) -> MonetaryAmount:
    if _is_missing(value):
        return MonetaryAmount(kind=kind, status=ValueStatus.UNKNOWN)
    source_text = str(value).strip()
    normalized = source_text.replace("R$", "").replace(" ", "").replace(".", "").replace(",", ".")
    try:
        amount = Decimal(normalized)
    except InvalidOperation:
        issues.append(
            ValidationIssue(
                "source.invalid_money",
                path,
                "valor monetário brasileiro inválido",
                source_text,
                "R$1.234,56",
            )
        )
        return MonetaryAmount(kind=kind, status=ValueStatus.UNKNOWN, source_text=source_text)
    return MonetaryAmount(
        kind=kind,
        status=ValueStatus.KNOWN,
        amount=amount,
        currency="BRL",
        source_text=source_text,
    )


def _parse_prize_tier(
    indexed: Mapping[str, object],
    tier_key: str,
    hits: int,
    issues: list[ValidationIssue],
) -> PrizeTierResult:
    winners = _parse_integer(
        _lookup(indexed, f"ganhadores{hits}acertos"),
        f"draw.prize_tiers.{tier_key}.winners",
        issues,
    )
    prize = _parse_money(
        _lookup(indexed, f"rateio{hits}acertos"),
        "paid_per_winner",
        f"draw.prize_tiers.{tier_key}.prize_per_winner",
        issues,
    )
    return PrizeTierResult(
        tier_key=tier_key,
        status=ValueStatus.KNOWN if winners is not None else ValueStatus.UNKNOWN,
        winners=winners,
        prize_per_winner=prize,
    )
