"""Provas offline de fontes complementares reais, sem adaptador de produção."""
import csv
from datetime import date
import hashlib
import io
import json
from pathlib import Path
import re
import zipfile

import pytest
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
REPORT = json.loads((ROOT / "docs/data/reanalise-etapa-2-atualizacao-completa.json").read_text(encoding="utf-8"))


def sample(name):
    item = next(s for s in REPORT["samples"] if s["name"] == name)
    return ROOT / item["path"]


@pytest.mark.parametrize("item", REPORT["samples"], ids=lambda s: s["name"])
def test_new_evidence_bytes(item):
    body = (ROOT / item["path"]).read_bytes()
    assert len(body) == item["byte_count"]
    assert hashlib.sha256(body).hexdigest() == item["sha256"]


@pytest.mark.parametrize("game,file,total", [("pb", "texas_powerball.csv", 131), ("mm", "texas_megamillions.csv", 87)])
def test_new_york_recent_primary_window_matches_already_preserved_texas(game, file, total):
    rows = json.loads(sample(f"ny_{game}_recent_rows").read_bytes())
    texas = csv.reader(io.StringIO((ROOT / "docs/data/fontes-atuais-2026-10-27/amostras" / file).read_text(encoding="utf-8-sig")))
    texas = {date(int(r[3]), int(r[1]), int(r[2])).isoformat(): r for r in texas}
    assert len(rows) == len({r["draw_date"] for r in rows}) == total
    summary = json.loads(sample(f"ny_{game}_recent_summary").read_bytes())[0]
    assert int(summary["count"]) == total
    for row in rows:
        numbers = list(map(int, row["winning_numbers"].split()))
        main = numbers[:5]
        special = numbers[5] if game == "pb" else int(row["mega_ball"])
        counterpart = texas[row["draw_date"][:10]]
        assert len(set(main)) == 5
        assert sorted(main) == sorted(map(int, counterpart[4:9]))
        assert special == int(counterpart[9])
        if game == "pb":
            assert int(row["multiplier"]) == int(counterpart[10])


def test_double_play_availability_is_partial_not_the_primary_coverage():
    rows = json.loads(sample("ny_pb_recent_rows").read_bytes())
    available = [r for r in rows if r.get("double_play_winning_numbers")]
    assert len(available) == 63 and len(rows) - len(available) == 68
    assert min(r["draw_date"][:10] for r in available) == "2026-06-03"
    for row in available:
        numbers = list(map(int, row["double_play_winning_numbers"].split()))
        assert len(numbers) == 6 and len(set(numbers[:5])) == 5
        assert all(1 <= n <= 69 for n in numbers[:5]) and 1 <= numbers[5] <= 26


@pytest.mark.parametrize("game,expected_draws", [("pb", 6), ("mm", 4)])
def test_colorado_real_export_scopes_and_recent_combination(game, expected_draws):
    book = load_workbook(sample(f"colorado_{game}_xlsx"), read_only=True, data_only=True)
    try:
        sheet = book.worksheets[0]
        sheet.reset_dimensions()
        rows = list(sheet.iter_rows(values_only=True))
    finally:
        book.close()
    assert len(rows) == expected_draws + 1
    latest = dict(zip(rows[0], rows[1]))
    ny = json.loads(sample(f"ny_{game}_recent_rows").read_bytes())[-1]
    assert list(map(int, latest["Winning Numbers"].split(" - "))) == list(map(int, ny["winning_numbers"].split()))[:5]
    assert latest["Jackpot"] > 0 and latest["Jackpot Cash Value"] > 0
    assert "Jackpot CO Winners" in latest
    if game == "pb":
        assert latest["Double Play Winning Numbers"] == "5 - 11 - 14 - 24 - 59"
        assert latest["Double Play Powerball"] == 8


def test_austrian_date_blocks_are_draws_not_tier_rows_and_match_fdj():
    body = sample("austria_2026").read_bytes()
    with pytest.raises(UnicodeDecodeError):
        body.decode("utf-8-sig")
    rows = list(csv.reader(io.StringIO(body.decode("cp1252")), delimiter=";"))
    draws = []
    for row in rows[1:]:
        if row[0].strip():
            day, month, year = map(int, re.search(r"(\d{2}\.\d{2}\.\d{4})", row[0]).group(1).split("."))
            draws.append({"date": date(year, month, day).isoformat(), "main": sorted(map(int, row[1:6])),
                          "stars": sorted(map(int, row[6:8])), "tiers": []})
        draws[-1]["tiers"].append(row[8:14])
    assert len(rows) == 1119 and len(draws) == len({d["date"] for d in draws}) == 86
    assert all(len(d["tiers"]) == 13 for d in draws)
    with zipfile.ZipFile(ROOT / "docs/data/fechamento-etapa-2-2026-10-27/amostras/fdj_recente.zip") as z:
        fdj = list(csv.DictReader(io.StringIO(z.read(z.namelist()[0]).decode("cp1252")), delimiter=";"))
    fdj = {"-".join(reversed(r["date_de_tirage"].split("/"))): r for r in fdj}
    for draw in draws:
        counterpart = fdj[draw["date"]]
        assert draw["main"] == sorted(int(counterpart[f"boule_{i}"]) for i in range(1, 6))
        assert draw["stars"] == sorted(int(counterpart[f"etoile_{i}"]) for i in range(1, 3))
    assert draws[-1]["tiers"][0][3] == "JP"
    assert "zugeschlagen" in draws[-1]["tiers"][1][5]
