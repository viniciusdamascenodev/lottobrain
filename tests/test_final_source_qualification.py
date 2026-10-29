"""Qualificação offline de evidência real; nenhum coletor ou parser de produção novo."""
import csv
from datetime import datetime
from decimal import Decimal
import hashlib
import io
import json
from pathlib import Path
import re
import zipfile

from openpyxl import load_workbook
import pytest

from app.domain import RawProvenance
from app.ingestion.mega_sena_api import EXPECTED_FIELDS, SOURCE_KEY, normalize_response

ROOT = Path(__file__).resolve().parents[1]
REPORT = json.loads((ROOT / "docs/data/qualificacao-final-etapa-2.json").read_bytes())
PREVIOUS = json.loads((ROOT / "docs/data/reanalise-etapa-2-atualizacao-completa.json").read_bytes())


def sample(name):
    return ROOT / next(s["path"] for s in REPORT["samples"] if s["name"] == name)


def previous_sample(name):
    return ROOT / next(s["path"] for s in PREVIOUS["samples"] if s["name"] == name)


@pytest.mark.parametrize("item", REPORT["samples"], ids=lambda s: s["name"])
def test_retained_evidence_integrity(item):
    body = (ROOT / item["path"]).read_bytes()
    assert len(body) == item["byte_count"]
    assert hashlib.sha256(body).hexdigest() == item["sha256"]


@pytest.mark.parametrize("name,contest,title", [
    ("caixa_mega_json", 3075, None),
    ("caixa_mega_special", 3010, "Mega-Sena 30 anos"),
])
def test_real_nullable_and_text_title_are_the_only_extra_fields(name, contest, title):
    body = sample(name).read_bytes()
    payload = json.loads(body)
    assert set(payload) == EXPECTED_FIELDS | {"tituloConcursoEspecial"}
    assert payload["tituloConcursoEspecial"] == title
    assert payload["numero"] == contest
    receipt = next(p for p in REPORT["probes"] if p["name"] == name)
    provenance = RawProvenance(SOURCE_KEY, receipt["sha256"],
                              datetime.fromisoformat(receipt["observed_at_utc"]), "qualification-only")
    # O normalizador de produção conserva sua recusa de schema novo.
    result = normalize_response(body, provenance, expected_contest=contest)
    assert not result.is_valid
    assert {i.code for i in result.issues} == {"source.schema_changed"}
    # Projeção exclusivamente de inspeção: a parte já contratada continua válida.
    # O título permanece no raw; isto não instala um adaptador que o descarte.
    known_fields = {k: v for k, v in payload.items() if k in EXPECTED_FIELDS}
    assert normalize_response(json.dumps(known_fields).encode(), provenance,
                              expected_contest=contest).is_valid


def test_transport_controls_preserve_secure_tls_and_do_not_claim_deployment_proof():
    probes = {p["name"]: p for p in REPORT["probes"]}
    assert len(probes) == 14
    for p in probes.values():
        assert p["http_status"] == 200 and p["tls_verified"] is True
        if "verify_mode" in p:
            assert p["verify_mode"] == 2 and p["check_hostname"] is True
    for game in ("pb", "mm"):
        regular = probes[f"control_ny_{game}_direct"]
        native = probes[f"ny_{game}_python_native"]
        assert regular["sha256"] == native["sha256"]
        assert regular["proxy_used"] is False and native["proxy_used"] is False
    for name in ("caixa_mega_download", "caixa_loto_download", "caixa_mili_download"):
        receipt = probes[name]
        assert any(s["reused_existing_file"] and s["sha256"] == receipt["sha256"] for s in REPORT["samples"])
    actual = probes["actual_existing_mega_client"]
    assert actual["normalizer_is_valid"] is True
    assert actual["operational_ingestion_called"] is False
    assert REPORT["all_sources_operational"] is False


def test_all_68_double_play_gaps_have_real_source_rows_and_overlap_agrees():
    ny = json.loads(previous_sample("ny_pb_recent_rows").read_bytes())
    ny = {r["draw_date"][:10]: r for r in ny}
    gaps = {d for d, r in ny.items() if not r.get("double_play_winning_numbers")}
    book = load_workbook(sample("colorado_all_xlsx"), read_only=True, data_only=True)
    checked = {}; overlap = 0; total = 0
    try:
        sheet = book.worksheets[0]
        sheet.reset_dimensions()
        rows = sheet.iter_rows(values_only=True)
        header = next(rows)
        assert len(header) == 61
        for ordinal, row in enumerate(rows, 2):
            total += 1
            match = re.search(r"\d{1,2}/\d{1,2}/\d{2,4}$", str(row[0]))
            if not match:
                continue
            text = match.group()
            day = datetime.strptime(text, "%m/%d/%y" if len(text.split("/")[-1]) == 2 else "%m/%d/%Y").date().isoformat()
            if day not in ny:
                continue
            assert day not in checked
            data = dict(zip(header, row)); ref = ny[day]
            primary = list(map(int, ref["winning_numbers"].split()))
            assert sorted(map(int, data["Winning Numbers"].split(" - "))) == sorted(primary[:5])
            assert int(data["Powerball"]) == primary[5]
            assert int(data["Power Play"]) == int(ref["multiplier"])
            dp = sorted(map(int, data["Double Play Winning Numbers"].split(" - ")))
            ball = int(data["Double Play Powerball"])
            assert len(dp) == len(set(dp)) == 5
            assert all(1 <= n <= 69 for n in dp) and 1 <= ball <= 26
            if ref.get("double_play_winning_numbers"):
                old = list(map(int, ref["double_play_winning_numbers"].split()))
                assert dp == sorted(old[:5]) and ball == old[5]
                overlap += 1
            checked[day] = {"date": day, "main": dp, "powerball": ball,
                            "source_row": ordinal, "filled_ny_gap": day in gaps}
    finally:
        book.close()
    assert total == 2904
    assert len(checked) == 131 and len(gaps) == 68 and overlap == 63
    assert not gaps - checked.keys()
    assert sorted(checked.values(), key=lambda r: r["date"]) == REPORT["double_play"]["recent_double_play"]
    assert REPORT["double_play"]["all_history_not_audited"] is True
    colorado = next(s for s in REPORT["source_use"] if s["provider"] == "Colorado Lottery")
    assert colorado["status"] == "automated_use_restricted"


def test_europe_winner_counts_and_financial_markers_are_not_flattened():
    body = previous_sample("austria_2026").read_bytes()
    with pytest.raises(UnicodeDecodeError):
        body.decode("utf-8-sig")
    rows = list(csv.reader(io.StringIO(body.decode("cp1252")), delimiter=";"))
    with zipfile.ZipFile(ROOT / "docs/data/fechamento-etapa-2-2026-10-27/amostras/fdj_recente.zip") as z:
        fdj = list(csv.DictReader(io.StringIO(z.read(z.namelist()[0]).decode("cp1252")), delimiter=";"))
    fdj = {datetime.strptime(r["date_de_tirage"], "%d/%m/%Y").date().isoformat(): r for r in fdj}
    day = None; draws = set(); tiers = 0; jackpots = 0; transfers = 0; numeric = 0
    for row in rows[1:]:
        if row[0].strip():
            day = datetime.strptime(re.search(r"\d{2}\.\d{2}\.\d{4}", row[0]).group(), "%d.%m.%Y").date().isoformat()
            draws.add(day)
        rank = int(row[8]); ref = fdj[day]; tiers += 1
        expected_winners = int(ref[f"nombre_de_gagnant_au_rang{rank}_Euro_Millions_en_europe"])
        text = row[11].strip(); quota = row[13].strip()
        fdj_payout = Decimal(ref[f"rapport_du_rang{rank}_Euro_Millions"].replace(",", "."))
        if text == "JP":
            assert rank == 1 and expected_winners == 0 and fdj_payout == 0
            assert Decimal(quota.replace(".", "").replace(",", ".")) > 0
            jackpots += 1
        else:
            assert int(text.replace(".", "")) == expected_winners
        if "zugeschlagen" in quota:
            assert quota == "wird dem Rang 5+0 zugeschlagen"
            assert rank == 2 and expected_winners == 0
            transfers += 1
        else:
            amount = Decimal(quota.replace(".", "").replace(",", "."))
            numeric += 1
            if text != "JP":
                assert amount == fdj_payout
    assert (len(draws), tiers, jackpots, transfers, numeric) == (86, 1118, 77, 10, 1108)
    assert REPORT["euro_millions_financial_comparison"]["transferred_amount"] is None
    assert REPORT["euro_millions_financial_comparison"]["actual_paid_amount_proven"] is False


def test_original_raw_baselines_are_preserved():
    assert len(REPORT["original_raw_sha256"]) == 6
    for file, expected in REPORT["original_raw_sha256"].items():
        assert hashlib.sha256((ROOT / file).read_bytes()).hexdigest() == expected
